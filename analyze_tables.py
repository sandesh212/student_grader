#!/usr/bin/env python3
"""Deep dive into table structure"""
from docx import Document

def analyze_tables(filepath):
    doc = Document(filepath)
    
    print(f"\n{'='*80}")
    print(f"TABLE CONTENT ANALYSIS")
    print(f"{'='*80}\n")
    
    for table_idx, table in enumerate(doc.tables, 1):
        print(f"\nTABLE {table_idx}:")
        print(f"  Dimensions: {len(table.rows)} rows × {len(table.columns)} columns")
        
        # Sample cells with content
        content_count = 0
        for row_idx, row in enumerate(table.rows):
            for cell_idx, cell in enumerate(row.cells):
                text = cell.text.strip()
                if text and len(text) > 20:
                    content_count += 1
                    if content_count <= 5:  # Show first 5 meaningful cells
                        # Check for red text in this cell
                        has_red = False
                        red_parts = []
                        black_parts = []
                        
                        for para in cell.paragraphs:
                            for run in para.runs:
                                run_text = run.text.strip()
                                if run_text:
                                    if run.font.color and run.font.color.rgb:
                                        r, g, b = run.font.color.rgb
                                        if r > 150 and g < 100 and b < 100:
                                            has_red = True
                                            red_parts.append(run_text)
                                        else:
                                            black_parts.append(run_text)
                                    else:
                                        black_parts.append(run_text)
                        
                        color_info = ""
                        if has_red:
                            color_info = f"\n      RED: {' '.join(red_parts[:3])}..."
                        if black_parts:
                            color_info += f"\n      BLACK: {' '.join(black_parts[:3])}..."
                        
                        print(f"\n  Row {row_idx+1}, Cell {cell_idx+1}:")
                        print(f"    Text: {text[:120]}{color_info}")
        
        if content_count > 5:
            print(f"\n  ... and {content_count - 5} more cells with content")
        
        if table_idx >= 2:  # Limit to first 2 tables for readability
            print(f"\n... and {len(doc.tables) - 2} more tables")
            break

if __name__ == "__main__":
    analyze_tables("samples/Knowledge-Seamanship.docx")
