#!/usr/bin/env python3
"""Quick diagnostic to check document structure"""
from docx import Document
import sys

def check_document(filepath):
    doc = Document(filepath)
    
    print(f"\n{'='*80}")
    print(f"DOCUMENT ANALYSIS: {filepath}")
    print(f"{'='*80}\n")
    
    print(f"Total Paragraphs: {len(doc.paragraphs)}")
    print(f"Total Tables: {len(doc.tables)}\n")
    
    print("FIRST 10 PARAGRAPHS:")
    print("-" * 80)
    for i, para in enumerate(doc.paragraphs[:10], 1):
        text = para.text.strip()
        if text:
            # Check for colored text
            has_red = False
            for run in para.runs:
                if run.font.color and run.font.color.rgb:
                    r, g, b = run.font.color.rgb
                    if r > 150 and g < 100 and b < 100:
                        has_red = True
                        break
            
            color_marker = " [HAS RED TEXT]" if has_red else ""
            print(f"{i}. {text[:100]}{color_marker}")
    
    print("\n" + "-" * 80)
    print("SAMPLE PARAGRAPH WITH RUNS:")
    print("-" * 80)
    
    # Find a paragraph with multiple runs
    for para in doc.paragraphs[:20]:
        if len(para.runs) > 1 and para.text.strip():
            print(f"Paragraph: {para.text[:100]}")
            print(f"Number of runs: {len(para.runs)}\n")
            for j, run in enumerate(para.runs[:5], 1):
                color_info = "DEFAULT"
                if run.font.color and run.font.color.rgb:
                    r, g, b = run.font.color.rgb
                    color_info = f"RGB({r},{g},{b})"
                print(f"  Run {j}: '{run.text[:50]}' - Color: {color_info}")
            break
    
    if doc.tables:
        print("\n" + "-" * 80)
        print("FIRST TABLE SAMPLE:")
        print("-" * 80)
        table = doc.tables[0]
        print(f"Rows: {len(table.rows)}, Columns: {len(table.columns)}\n")
        for i, row in enumerate(table.rows[:3], 1):
            print(f"Row {i}:")
            for j, cell in enumerate(row.cells[:3], 1):
                print(f"  Cell {j}: {cell.text[:60]}")

if __name__ == "__main__":
    if len(sys.argv) > 1:
        check_document(sys.argv[1])
    else:
        check_document("samples/Knowledge-Seamanship.docx")
