import os
import re
from docx import Document

# ---------- UNIT CACHE LOADER ----------

def load_all_units(cache_dir="unit_cache"):
    """Load all unit information from cached files."""
    units = {}
    if not os.path.exists(cache_dir):
        print(f"WARNING: {cache_dir} directory not found.")
        return units
    
    for filename in os.listdir(cache_dir):
        if filename.endswith('.txt'):
            unit_code = filename.replace('.txt', '')
            filepath = os.path.join(cache_dir, filename)
            try:
                with open(filepath, 'r', encoding='utf-8') as f:
                    content = f.read()
                    units[unit_code] = {
                        'code': unit_code,
                        'title': extract_unit_title(content),
                        'full_text': content.lower()
                    }
            except Exception as e:
                print(f"WARNING: Could not load {filename}: {e}")
    
    print(f"Loaded {len(units)} units from cache.")
    return units

def extract_unit_title(content):
    """Try to extract unit title from various formats."""
    # Try to find title patterns
    patterns = [
        r'"title":"([^"]+)"',
        r'<title>([^<]+)</title>',
        r'Unit of Competency[:\s]+([^\n]+)',
        r'([A-Z][a-z]+(?:\s+[a-z]+){2,})'
    ]
    for pattern in patterns:
        match = re.search(pattern, content)
        if match:
            return match.group(1).strip()
    return ''

# ---------- QUESTION & ANSWER EXTRACTION ----------

def is_red_text(run):
    """Check if text is red colored (answers in marking sheets)."""
    if run.font.color and run.font.color.rgb:
        r, g, b = run.font.color.rgb
        return r > 150 and g < 100 and b < 100
    return False

def extract_from_document(doc):
    """Extract questions and answers from a document."""
    results = []
    
    for table in doc.tables:
        for row in table.rows:
            if len(row.cells) == 0:
                continue
            
            question_cell = row.cells[0]
            question_text = question_cell.text.strip()
            
            # Skip header rows, empty cells, and short section titles
            if not question_text or len(question_text) < 20:
                continue
            
            # Skip common headers
            skip_headers = ['S', 'NS', 'Emergency Preparedness', 'Vessel Handling', 
                          'Participant Name', 'RTO Name', 'Course Name', 'UoC']
            if any(header in question_text for header in skip_headers):
                continue
            
            # Check if this looks like a question
            has_question_words = bool(re.search(
                r'\b(list|describe|name|state|explain|identify|give|provide|what|how|why|which|match)\b', 
                question_text.lower()
            ))
            has_question_mark = '?' in question_text
            has_blanks = '_' in question_text
            
            if has_question_words or has_question_mark or has_blanks:
                # Extract answers from colored text (red = answers)
                answers = []
                for para in question_cell.paragraphs:
                    for run in para.runs:
                        if run.text.strip() and is_red_text(run):
                            answer_text = run.text.strip()
                            # Split multi-part answers
                            parts = re.split(r'\n|;(?!\w)', answer_text)
                            for part in parts:
                                part = part.strip()
                                if part and len(part) > 2 and not part.startswith('_'):
                                    answers.append(part)
                
                # Clean up the question text
                question_clean = re.sub(r'_{3,}', '[___]', question_text)
                question_clean = re.sub(r'\s+', ' ', question_clean).strip()
                
                # Extract unit codes (e.g., K007-PC1.1, N008:K9)
                unit_codes = extract_unit_codes(question_text)
                
                results.append({
                    'question': question_clean,
                    'answers': answers,
                    'unit_codes': unit_codes
                })
    
    return results

def extract_unit_codes(text):
    """Extract unit codes from question text (e.g., K007-PC1.1 -> MARK007)."""
    # Find patterns like K007, N008, I003, C037, etc.
    codes = re.findall(r'\b([A-Z]\d{3})[:\-]', text)
    
    # Convert to full unit codes (K007 -> MARK007)
    full_codes = []
    for code in codes:
        full_code = "MAR" + code
        full_codes.append(full_code)
    
    return list(set(full_codes))  # Remove duplicates

# ---------- UNIT MATCHING ----------

def match_to_units(question_data, all_units):
    """Match a question to relevant units."""
    matches = []
    question_text = question_data['question']
    unit_codes = question_data['unit_codes']
    
    # Combine question and answers for keyword matching
    full_context = (question_text + ' ' + ' '.join(question_data['answers'])).lower()
    key_terms = set(re.findall(r'\b[a-z]{4,}\b', full_context))
    
    for unit_code, unit_data in all_units.items():
        score = 0
        
        # HIGH PRIORITY: Explicit unit code mention
        if unit_code in unit_codes:
            score += 200
            matches.append({
                'unit_code': unit_code,
                'unit_title': unit_data['title'],
                'score': score,
                'match_type': 'Explicit reference'
            })
            continue
        
        # LOWER PRIORITY: Keyword matching for non-explicit matches
        if unit_code.lower() in full_context:
            score += 100
        
        title_words = set(re.findall(r'\b[a-z]{4,}\b', unit_data['title'].lower()))
        common_words = key_terms & title_words
        score += len(common_words) * 10
        
        # Only add if there's a reasonable match
        if score > 30:
            matches.append({
                'unit_code': unit_code,
                'unit_title': unit_data['title'],
                'score': score,
                'match_type': 'Keyword match'
            })
    
    # Sort by score
    matches.sort(key=lambda x: x['score'], reverse=True)
    return matches[:3]  # Top 3 matches

# ---------- MAIN PROCESSING ----------

def process_pair(question_file, answer_file, all_units):
    """Process a question file and its corresponding answer/marking sheet file."""
    print(f"\n{'='*100}")
    print(f"PROCESSING: {os.path.basename(question_file)}")
    print(f"{'='*100}\n")
    
    # Load question document
    try:
        q_doc = Document(question_file)
        questions = extract_from_document(q_doc)
    except Exception as e:
        print(f"ERROR: Could not process {question_file}: {e}")
        return
    
    # Load answer document if it exists
    answers_from_marking = {}
    if answer_file and os.path.exists(answer_file):
        try:
            a_doc = Document(answer_file)
            answer_data = extract_from_document(a_doc)
            # Create a mapping by question number
            for i, item in enumerate(answer_data, 1):
                if item['answers']:
                    answers_from_marking[i] = item['answers']
        except Exception as e:
            print(f"WARNING: Could not process {answer_file}: {e}")
    
    # Display results
    for i, q_item in enumerate(questions, 1):
        print(f"\n{'─'*100}")
        print(f"Q{i}. {q_item['question'][:150]}")
        
        # Check if we have answers from the question doc or marking sheet
        answers = q_item['answers'] or answers_from_marking.get(i, [])
        
        if answers:
            print(f"\n   ✓ ANSWERS ({len(answers)}):")
            for j, ans in enumerate(answers, 1):
                print(f"     {j}) {ans[:120]}")
        else:
            print(f"\n   ⚠️  No answers found")
        
        # Match to units
        if q_item['unit_codes']:
            print(f"\n   📚 UNIT CODES DETECTED: {', '.join(q_item['unit_codes'])}")
        
        matches = match_to_units(q_item, all_units)
        if matches:
            print(f"\n   🎯 MATCHED UNITS:")
            for match in matches:
                print(f"      • {match['unit_code']}: {match['unit_title'][:60]}")
                print(f"        ({match['match_type']}, Score: {match['score']})")
        else:
            print(f"\n   ⚠️  No unit matches found")
    
    print(f"\n{'='*100}")
    print(f"SUMMARY: {len(questions)} questions processed from {os.path.basename(question_file)}")
    print(f"{'='*100}\n")

# ---------- MAIN EXECUTION ----------

def main():
    # Load all units first
    all_units = load_all_units()
    
    if not all_units:
        print("ERROR: No units loaded. Cannot proceed.")
        return
    
    # Get input file
    filename = input("Enter DOCX filename (or press Enter to process all samples): ").strip()
    
    if not filename:
        # Process all sample files
        samples_dir = "samples"
        if not os.path.exists(samples_dir):
            print(f"ERROR: {samples_dir} directory not found.")
            return
        
        files = [f for f in os.listdir(samples_dir) if f.endswith('.docx') and not f.startswith('~$')]
        
        # Group files into question/answer pairs
        processed = set()
        for file in files:
            if file in processed:
                continue
            
            if 'Marking' in file or 'marking' in file:
                continue  # Skip marking sheets in the main loop
            
            question_path = os.path.join(samples_dir, file)
            
            # Try to find corresponding marking sheet
            base_name = file.replace('.docx', '')
            possible_answer_names = [
                f"{base_name}-Marking-Sheet.docx",
                f"{base_name} Marking Sheet.docx",
                f"{base_name}-marking-sheet.docx"
            ]
            
            answer_path = None
            for ans_name in possible_answer_names:
                ans_path = os.path.join(samples_dir, ans_name)
                if os.path.exists(ans_path):
                    answer_path = ans_path
                    processed.add(ans_name)
                    break
            
            process_pair(question_path, answer_path, all_units)
            processed.add(file)
    else:
        # Process single file
        path_in_cwd = os.path.join(os.getcwd(), filename)
        path_in_samples = os.path.join(os.getcwd(), "samples", filename)
        
        if os.path.isfile(path_in_cwd):
            path = path_in_cwd
        elif os.path.isfile(path_in_samples):
            path = path_in_samples
        else:
            print(f"ERROR: File '{filename}' not found.")
            return
        
        # Try to find marking sheet
        base_name = os.path.basename(path).replace('.docx', '')
        dir_name = os.path.dirname(path)
        answer_path = None
        
        for suffix in ['-Marking-Sheet', ' Marking Sheet', '-marking-sheet']:
            ans_path = os.path.join(dir_name, f"{base_name}{suffix}.docx")
            if os.path.exists(ans_path):
                answer_path = ans_path
                break
        
        process_pair(path, answer_path, all_units)

if __name__ == "__main__":
    main()
